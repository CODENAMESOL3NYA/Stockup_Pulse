"""Strict, local CSV/XLSX import for sales history and current stock counts."""

import csv
import io
import math
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Any, Dict, Iterable, Tuple

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

XLSX_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
XLSX_MAX_MEMBERS = 256
XLSX_MAX_SHEETS = 10
XLSX_MAX_ROWS = 100_000
XLSX_MAX_COLUMNS = 100
XLSX_MAX_CELLS = 1_000_000


class ImportValidationError(ValueError):
    """Raised when an uploaded sales or stock file violates the import contract."""


def normalize_product_name(value: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", value.casefold()))


def _number(value: Any, *, field: str, row: int) -> float:
    if isinstance(value, bool):
        raise ImportValidationError(f"Row {row}: {field} must be a finite nonnegative number.")
    try:
        parsed = float(value)
    except (TypeError, ValueError) as exc:
        raise ImportValidationError(f"Row {row}: {field} must be a finite nonnegative number.") from exc
    if not math.isfinite(parsed) or parsed < 0:
        raise ImportValidationError(f"Row {row}: {field} must be a finite nonnegative number.")
    return parsed


def _records_from_csv(content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportValidationError("CSV files must use UTF-8 encoding.") from exc
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise ImportValidationError("CSV file is empty or has no header row.")
    normalized_headers = [
        str(name).strip().lower() if name is not None else ""
        for name in reader.fieldnames
    ]
    if not all(normalized_headers) or len(set(normalized_headers)) != len(normalized_headers):
        raise ImportValidationError("CSV header names must be present and unique.")
    headers = set(normalized_headers)
    rows: list[dict[str, Any]] = []
    for row in reader:
        if None in row:
            raise ImportValidationError("CSV rows must not contain more fields than the header.")
        rows.append({str(key).strip().lower(): value for key, value in row.items() if key is not None})
    return [{"headers": headers, "rows": rows}]


def _records_from_xlsx(content: bytes) -> dict[str, list[dict[str, Any]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > XLSX_MAX_MEMBERS:
                raise ImportValidationError("XLSX contains too many internal files.")
            total_uncompressed = sum(member.file_size for member in members)
            if total_uncompressed > XLSX_MAX_UNCOMPRESSED_BYTES:
                raise ImportValidationError("XLSX expands beyond the 100 MiB processing limit.")
            if any(member.flag_bits & 0x1 for member in members):
                raise ImportValidationError("Encrypted XLSX files are not supported.")
            if any(
                member.file_size > 0
                and (member.compress_size == 0 or member.file_size / member.compress_size > 1000)
                for member in members
            ):
                raise ImportValidationError("XLSX contains an unsafe compression ratio.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ImportValidationError:
        raise
    except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise ImportValidationError("XLSX file could not be read.") from exc
    result: dict[str, list[dict[str, Any]]] = {}
    try:
        if len(workbook.sheetnames) > XLSX_MAX_SHEETS:
            raise ImportValidationError("XLSX contains too many worksheets.")
        for sheet_name in ("sales", "stock"):
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            if (
                (sheet.max_row is not None and sheet.max_row > XLSX_MAX_ROWS)
                or (sheet.max_column is not None and sheet.max_column > XLSX_MAX_COLUMNS)
            ):
                raise ImportValidationError(
                    f"XLSX sheet '{sheet_name}' exceeds the row or column processing limit."
                )
            values = sheet.iter_rows(values_only=True)
            header_row = next(values, None)
            if not header_row:
                raise ImportValidationError(f"XLSX sheet '{sheet_name}' is empty.")
            if len(header_row) > XLSX_MAX_COLUMNS:
                raise ImportValidationError(f"XLSX sheet '{sheet_name}' has too many columns.")
            headers = [str(value).strip().lower() if value is not None else "" for value in header_row]
            if not all(headers) or len(set(headers)) != len(headers):
                raise ImportValidationError(f"XLSX sheet '{sheet_name}' headers must be present and unique.")
            data_rows = []
            parsed_cells = len(header_row)
            for row_number, cells in enumerate(values, start=2):
                if row_number > XLSX_MAX_ROWS:
                    raise ImportValidationError(f"XLSX sheet '{sheet_name}' has too many rows.")
                if len(cells) > XLSX_MAX_COLUMNS:
                    raise ImportValidationError(f"XLSX sheet '{sheet_name}' has too many columns.")
                if not any(value is not None and str(value).strip() for value in cells):
                    continue
                parsed_cells += len(cells)
                if parsed_cells > XLSX_MAX_CELLS:
                    raise ImportValidationError(f"XLSX sheet '{sheet_name}' has too many cells.")
                data_rows.append({
                    key: cells[index] if index < len(cells) else None
                    for index, key in enumerate(headers)
                    if key
                })
            result[sheet_name] = [{"headers": set(headers), "rows": data_rows}]
    finally:
        workbook.close()
    if not result:
        raise ImportValidationError("XLSX must contain a sheet named 'sales' or 'stock'.")
    return result


def parse_sales_stock_files(files: Iterable[Tuple[str, bytes]]) -> Dict[str, Any]:
    """Parse and combine supported tabular inputs with deterministic aggregation."""
    sales_rows: list[dict[str, Any]] = []
    stock_rows: list[dict[str, Any]] = []
    stock_names: dict[str, str] = {}

    def consume(kind: str, records: list[dict[str, Any]]) -> None:
        expected = {"date", "product", "quantity_sold"} if kind == "sales" else {"product", "current_stock"}
        for record_set in records:
            headers = record_set["headers"]
            missing = expected - headers
            unexpected = headers - expected
            if missing or unexpected:
                raise ImportValidationError(
                    f"{kind.title()} data must contain exactly these columns: {', '.join(sorted(expected))}."
                )
            for row_number, row in enumerate(record_set["rows"], start=2):
                product = str(row.get("product") or "").strip()
                if not product:
                    raise ImportValidationError(f"Row {row_number}: product is required.")
                key = normalize_product_name(product)
                if not key:
                    raise ImportValidationError(f"Row {row_number}: product name has no searchable characters.")
                if kind == "sales":
                    raw_date = row.get("date")
                    if isinstance(raw_date, datetime):
                        sale_date = raw_date.date().isoformat()
                    elif isinstance(raw_date, date):
                        sale_date = raw_date.isoformat()
                    else:
                        raw_date = str(raw_date or "").strip()
                        try:
                            sale_date = date.fromisoformat(raw_date).isoformat()
                        except ValueError as exc:
                            raise ImportValidationError(
                                f"Row {row_number}: date must use ISO YYYY-MM-DD format."
                            ) from exc
                        if raw_date != sale_date:
                            raise ImportValidationError(
                                f"Row {row_number}: date must use ISO YYYY-MM-DD format."
                            )
                    sales_rows.append({
                        "key": key,
                        "product": product,
                        "date": sale_date,
                        "quantity": _number(row.get("quantity_sold"), field="quantity_sold", row=row_number),
                    })
                else:
                    if key in stock_names:
                        raise ImportValidationError(
                            f"Duplicate stock rows found for '{stock_names[key]}' and '{product}'."
                        )
                    stock_names[key] = product
                    stock_rows.append({
                        "key": key,
                        "product": product,
                        "quantity": _number(row.get("current_stock"), field="current_stock", row=row_number),
                    })

    for filename, content in files:
        suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
        if suffix == "csv":
            records = _records_from_csv(content)
            headers = records[0]["headers"]
            sales_signature = {"date", "product", "quantity_sold"}
            stock_signature = {"product", "current_stock"}
            if sales_signature.issubset(headers):
                consume("sales", records)
            elif stock_signature.issubset(headers):
                consume("stock", records)
            else:
                raise ImportValidationError(
                    "CSV headers must be date, product, quantity_sold or product, current_stock."
                )
        elif suffix == "xlsx":
            workbook_records = _records_from_xlsx(content)
            if "sales" in workbook_records:
                consume("sales", workbook_records["sales"])
            if "stock" in workbook_records:
                consume("stock", workbook_records["stock"])
        else:
            raise ImportValidationError("Only CSV and XLSX files can be imported as sales or stock data.")

    if not sales_rows:
        raise ImportValidationError("At least one sales row is required.")
    if not stock_rows:
        raise ImportValidationError("At least one current-stock row is required.")

    names: dict[str, set[str]] = defaultdict(set)
    grouped_sales: dict[tuple[str, str], list[float]] = defaultdict(list)
    for row in sales_rows:
        names[row["key"]].add(row["product"])
        grouped_sales[(row["key"], row["date"])].append(row["quantity"])
    for row in stock_rows:
        names[row["key"]].add(row["product"])

    first_day = min(date.fromisoformat(row["date"]) for row in sales_rows)
    last_day = max(date.fromisoformat(row["date"]) for row in sales_rows)
    sales: dict[str, dict[str, float]] = {}
    for key in sorted({row["key"] for row in sales_rows}):
        daily: dict[str, float] = {}
        cursor = first_day
        while cursor <= last_day:
            day_key = cursor.isoformat()
            try:
                total = math.fsum(grouped_sales.get((key, day_key), []))
            except OverflowError as exc:
                raise ImportValidationError("Aggregated sales quantity exceeds the supported numeric range.") from exc
            if not math.isfinite(total):
                raise ImportValidationError("Aggregated sales quantity exceeds the supported numeric range.")
            daily[day_key] = total
            cursor += timedelta(days=1)
        sales[key] = daily

    return {
        "sales": sales,
        "stock": {row["key"]: row["quantity"] for row in stock_rows},
        "product_names": {key: sorted(values) for key, values in names.items()},
        "sales_interval": {"start": first_day.isoformat(), "end": last_day.isoformat()},
    }


def normalize_leaflet_candidates(
    candidates: list[dict[str, Any]],
    *,
    source: str,
    source_filename: str,
) -> list[dict[str, Any]]:
    """Validate provider output and derive per-unit prices only for explicit pack counts."""
    normalized = []
    for index, candidate in enumerate(candidates, start=1):
        product = str(candidate.get("product") or candidate.get("product_name") or "").strip()
        pack_size = str(candidate.get("pack_size") or "").strip()
        supplier = str(candidate.get("supplier") or "").strip()
        if not product:
            raise ImportValidationError(f"Extracted item {index} must include a product name.")
        raw_price = candidate.get("pack_price", candidate.get("supplier_price", candidate.get("price")))
        price = None if raw_price is None else _number(raw_price, field="pack price", row=index)
        if price is not None and price <= 0:
            raise ImportValidationError(f"Extracted item {index} pack price must be greater than zero.")
        count_match = re.match(r"^\s*(\d+)\s*[xX×]\s*", pack_size)
        pack_count = int(count_match.group(1)) if count_match else None
        explicit_count = candidate.get("pack_count")
        if explicit_count is not None:
            if isinstance(explicit_count, bool):
                raise ImportValidationError(f"Extracted item {index} pack count must be a positive whole number.")
            try:
                parsed_count = float(explicit_count)
            except (TypeError, ValueError) as exc:
                raise ImportValidationError(f"Extracted item {index} pack count must be a positive whole number.") from exc
            if not math.isfinite(parsed_count) or not parsed_count.is_integer() or parsed_count < 1:
                raise ImportValidationError(f"Extracted item {index} pack count must be a positive whole number.")
            if pack_count is not None and pack_count != int(parsed_count):
                raise ImportValidationError(f"Extracted item {index} has conflicting pack-count values.")
            pack_count = int(parsed_count)
        unit_price = round(price / pack_count, 2) if pack_count and price is not None else None
        raw_confidence = candidate.get("confidence")
        confidence = 0.0
        if raw_confidence is not None:
            if isinstance(raw_confidence, bool):
                raise ImportValidationError(f"Extracted item {index} confidence must be between 0 and 1.")
            try:
                confidence = float(raw_confidence)
            except (TypeError, ValueError) as exc:
                raise ImportValidationError(f"Extracted item {index} confidence must be between 0 and 1.") from exc
            if not math.isfinite(confidence) or not 0 <= confidence <= 1:
                raise ImportValidationError(f"Extracted item {index} confidence must be between 0 and 1.")
        raw_warnings = candidate.get("warnings", [])
        if not isinstance(raw_warnings, list) or not all(isinstance(warning, str) for warning in raw_warnings):
            raise ImportValidationError(f"Extracted item {index} warnings must be a list of strings.")
        currency_value = str(candidate.get("currency") or "").strip().upper()
        if currency_value and not re.fullmatch(r"[A-Z]{3}", currency_value):
            raise ImportValidationError(f"Extracted item {index} currency must be a three-letter code.")
        warnings = list(raw_warnings)
        if not supplier:
            warnings.append("Supplier was not supplied; enter and verify the supplier.")
        if price is None:
            warnings.append("Supplier pack price was not supplied; enter the observed price.")
        if not currency_value:
            warnings.append("Currency was not supplied; verify the amount is in ZAR.")
        if not pack_count:
            warnings.append("Pack count is unclear; supplier unit price requires manual review.")
        normalized.append({
            "product": product,
            "pack_size": pack_size,
            "variant": str(candidate.get("variant") or candidate.get("pack_variant") or "").strip() or None,
            "pack_count": pack_count,
            "supplier_price": round(price, 2) if price is not None else None,
            "unit_price": unit_price,
            "our_price": None,
            "supplier": supplier,
            "currency": currency_value or None,
            "availability": str(candidate.get("availability") or "").strip() or None,
            "evidence": str(candidate.get("evidence") or candidate.get("offer_text") or "").strip() or None,
            "confidence": confidence,
            "source": source,
            "source_filename": source_filename,
            "warnings": warnings,
        })
    return normalized
